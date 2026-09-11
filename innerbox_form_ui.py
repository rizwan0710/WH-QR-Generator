# innerbox_form_ui.py - PART 1: IMPORTS & DUAL EXCEL ENGINE (APPEND MODE)
import tkinter as tk
from tkinter import messagebox, filedialog
from tkcalendar import DateEntry
import form_warehouse as backend
import custom_dropdown as cd  
import sqlite3
import os

try:
    import openpyxl
except ImportError:
    pass

def alih_fokus(event, input_seterusnya):
    input_seterusnya.focus_set()
    return "break"  

def laksanakan_import_excel_langsung(win_induk):
    """📥 Import Product Master (APPEND MODE - Keep old data & add new)"""
    path_fail = filedialog.askopenfilename(
        title="Select Customer Master List Excel File",
        filetypes=[("Excel Files", "*.xlsx *.xls"), ("All Files", "*.*")],
        parent=win_induk
    )
    if not path_fail:
        return
        
    conn = sqlite3.connect("warehouse_data.db")
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS master_produk (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            customer_name TEXT,
            drawing_no TEXT,
            part_number TEXT
        )
    """)
    
    query = "INSERT INTO master_produk (customer_name, drawing_no, part_number) VALUES (?, ?, ?)"
    try:
        wb = openpyxl.load_workbook(path_fail, data_only=True)
        sheet = wb.active
        kira = 0
        
        for row in sheet.iter_rows(min_row=2, max_col=4, values_only=True):
            if row and len(row) >= 4:
                cust_val = str(row[1]).strip().upper() if row[1] is not None else ""
                draw_val = str(row[2]).strip().upper() if row[2] is not None else ""
                part_val = str(row[3]).strip().upper() if row[3] is not None else ""
                
                if cust_val and cust_val != "NONE" and cust_val != "":
                    # 🔍 DUPLICATE CHECK
                    cursor.execute("""
                        SELECT id FROM master_produk 
                        WHERE customer_name = ? AND drawing_no = ? AND part_number = ?
                    """, (cust_val, draw_val, part_val))
                    
                    if not cursor.fetchone():
                        cursor.execute(query, (cust_val, draw_val, part_val))
                        kira += 1
                    
        conn.commit()
        messagebox.showinfo("SUCCESS", f"Successfully added {kira} new product master records!", parent=win_induk)
    except Exception as e:
        conn.rollback()
        messagebox.showerror("IMPORT ERROR", f"Failed to read Product Excel file:\n{str(e)}", parent=win_induk)
    finally:
        conn.close()

def laksanakan_import_excel_machine(win_induk):
    """📥 Import Machine Mapping (APPEND MODE - Keep old data & add new)"""
    path_fail = filedialog.askopenfilename(
        title="Select Machine Master Excel File",
        filetypes=[("Excel Files", "*.xlsx *.xls"), ("All Files", "*.*")],
        parent=win_induk
    )
    if not path_fail:
        return
        
    conn = sqlite3.connect("warehouse_data.db")
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS master_machine (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            part_number TEXT,
            machine_code TEXT
        )
    """)
    
    query = "INSERT INTO master_machine (part_number, machine_code) VALUES (?, ?)"
    try:
        wb = openpyxl.load_workbook(path_fail, data_only=True)
        sheet = wb.active
        
        header_row = [str(cell).strip().upper() for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        
        part_idx = None
        mach_idx = None
        
        for idx, header in enumerate(header_row):
            if "P/N" in header or "PART" in header:
                part_idx = idx
            if "M/C" in header or "MACHINE" in header or "MC" in header:
                mach_idx = idx
                
        if part_idx is None: part_idx = 2
        if mach_idx is None: mach_idx = 3
            
        kira = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row and len(row) > max(part_idx, mach_idx):
                part_val = str(row[part_idx]).strip().upper() if row[part_idx] is not None else ""
                mach_val = str(row[mach_idx]).strip().upper() if row[mach_idx] is not None else ""
                
                if part_val and mach_val and part_val != "NONE" and mach_val != "":
                    # 🔍 DUPLICATE CHECK
                    cursor.execute("SELECT id FROM master_machine WHERE part_number = ? AND machine_code = ?", (part_val, mach_val))
                    if not cursor.fetchone():
                        cursor.execute(query, (part_val, mach_val))
                        kira += 1
                    
        conn.commit()
        messagebox.showinfo("SUCCESS", f"Successfully added {kira} new Machine Code mapping records!", parent=win_induk)
    except Exception as e:
        conn.rollback()
        messagebox.showerror("IMPORT ERROR", f"Failed to read Machine Excel file:\n{str(e)}", parent=win_induk)
    finally:
        conn.close()
# innerbox_form_ui.py - PART 2: FORM INITIALIZATION & FIELDS 1-4
def buka_borang_warehouse(root):
    win_inner = tk.Toplevel(root)
    win_inner.title("FORM 1: INNER BOX PRODUCTION ENTRY")
    win_inner.geometry("540x660+450+30") 
    win_inner.configure(bg="#F8F9FA")
    win_inner.resizable(False, False)
    win_inner.grab_set()
    
    win_inner.bind("<Button-1>", lambda e: cd.tutup_dropdown() if e.widget.winfo_class() not in ["Entry", "Button"] else None)
    
    tk.Label(win_inner, text="INNER BOX PRODUCTION FORM", font=("Segoe UI", 13, "bold"), fg="#1E3A8A", bg="#F8F9FA").pack(pady=(15, 2))
    tk.Label(win_inner, text="ENTER PRODUCTION DETAILS PRECISELY:", font=("Segoe UI", 9, "italic"), fg="#64748B", bg="#F8F9FA").pack(pady=(0, 10))
    
    frame_data_entry = tk.LabelFrame(win_inner, text=" DATA ENTRY ", font=("Segoe UI", 8, "bold"), fg="#64748B", bg="#F8F9FA", padx=25, pady=10)
    frame_data_entry.pack(fill=tk.BOTH, expand=True, padx=30, pady=(0, 15))
    
    lbl_style = {"font": ("Segoe UI", 9, "bold"), "fg": "#334155", "bg": "#F8F9FA", "width": 20, "anchor": "w"}
    ent_style = {"font": ("Segoe UI", 10), "relief": "groove", "bd": 1}
    
    kamus_input = {}

    # Field 1: Date Created
    frame_row1 = tk.Frame(frame_data_entry, bg="#F8F9FA")
    frame_row1.pack(fill=tk.X, pady=4)
    tk.Label(frame_row1, text="Date Created :", **lbl_style).pack(side=tk.LEFT)
    kamus_input["date"] = DateEntry(frame_row1, width=30, font=("Segoe UI", 10), background="#0D6EFD", foreground="white", borderwidth=1, date_pattern="dd/mm/yyyy")
    kamus_input["date"].pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=3)
    
    def clear_isian_bawah(event=None):
        kamus_input["drawing"].delete(0, tk.END)
        kamus_input["part"].delete(0, tk.END)
        kamus_input["machine"].delete(0, tk.END)

    # Field 2: Customer Name
    frame_row2 = tk.Frame(frame_data_entry, bg="#F8F9FA")
    frame_row2.pack(fill=tk.X, pady=4)
    tk.Label(frame_row2, text="Customer Name :", **lbl_style).pack(side=tk.LEFT)
    frame_combobox_mix = tk.Frame(frame_row2, bg="#F8F9FA")
    frame_combobox_mix.pack(side=tk.LEFT, fill=tk.X, expand=True)
    kamus_input["customer"] = tk.Entry(frame_combobox_mix, font=("Segoe UI", 10), relief="groove", bd=1)
    kamus_input["customer"].pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=1)
    
    kamus_input["customer"].bind("<KeyRelease>", lambda e: cd.kendalikan_taipan_dropdown(
        kamus_input["customer"], win_inner, cd.ambil_senarai_customer_master()
    ))
    kamus_input["customer"].bind("<<Modified>>", clear_isian_bawah)
    
    tk.Button(frame_combobox_mix, text="▼", font=("Segoe UI", 7), bg="#E2E8F0", fg="#475569", relief="groove", bd=1, width=3, takefocus=False,
              command=lambda: [clear_isian_bawah(), cd.papar_senarai_toplevel(kamus_input["customer"], win_inner, cd.ambil_senarai_customer_master())]).pack(side=tk.RIGHT, fill=tk.Y)
    
    # Field 3: Drawing No
    frame_row3 = tk.Frame(frame_data_entry, bg="#F8F9FA")
    frame_row3.pack(fill=tk.X, pady=4)
    tk.Label(frame_row3, text="Drawing No :", **lbl_style).pack(side=tk.LEFT)
    frame_draw_mix = tk.Frame(frame_row3, bg="#F8F9FA")
    frame_draw_mix.pack(side=tk.LEFT, fill=tk.X, expand=True)
    kamus_input["drawing"] = tk.Entry(frame_draw_mix, font=("Segoe UI", 10), relief="groove", bd=1)
    kamus_input["drawing"].pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=1)
    
    kamus_input["drawing"].bind("<KeyRelease>", lambda e: cd.kendalikan_taipan_dropdown(
        kamus_input["drawing"], win_inner, cd.ambil_drawing_terikat(kamus_input["customer"].get())
    ))
    kamus_input["drawing"].bind("<<Modified>>", lambda e: [kamus_input["part"].delete(0, tk.END), kamus_input["machine"].delete(0, tk.END)])
    
    tk.Button(frame_draw_mix, text="▼", font=("Segoe UI", 7), bg="#E2E8F0", fg="#475569", relief="groove", bd=1, width=3, takefocus=False,
              command=lambda: [kamus_input["part"].delete(0, tk.END), kamus_input["machine"].delete(0, tk.END), cd.papar_senarai_toplevel(kamus_input["drawing"], win_inner, cd.ambil_drawing_terikat(kamus_input["customer"].get()))]).pack(side=tk.RIGHT, fill=tk.Y)

    # Field 4: Part Number
    frame_row4 = tk.Frame(frame_data_entry, bg="#F8F9FA")
    frame_row4.pack(fill=tk.X, pady=4)
    tk.Label(frame_row4, text="Part Number :", **lbl_style).pack(side=tk.LEFT)
    frame_part_mix = tk.Frame(frame_row4, bg="#F8F9FA")
    frame_part_mix.pack(side=tk.LEFT, fill=tk.X, expand=True)
    kamus_input["part"] = tk.Entry(frame_part_mix, font=("Segoe UI", 10), relief="groove", bd=1)
    kamus_input["part"].pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=1)
    
    kamus_input["part"].bind("<KeyRelease>", lambda e: cd.kendalikan_taipan_dropdown(
        kamus_input["part"], win_inner, cd.ambil_part_terikat(kamus_input["customer"].get(), kamus_input["drawing"].get())
    ))
    kamus_input["part"].bind("<<Modified>>", lambda e: kamus_input["machine"].delete(0, tk.END))
    
    tk.Button(frame_part_mix, text="▼", font=("Segoe UI", 7), bg="#E2E8F0", fg="#475569", relief="groove", bd=1, width=3, takefocus=False,
              command=lambda: [kamus_input["machine"].delete(0, tk.END), cd.papar_senarai_toplevel(kamus_input["part"], win_inner, cd.ambil_part_terikat(kamus_input["customer"].get(), kamus_input["drawing"].get()))]).pack(side=tk.RIGHT, fill=tk.Y)
# innerbox_form_ui.py - PART 3: BAKI FIELDS & SUBMIT (LIVE PRODUCTION ENGINE)
    # Field 5: Total Quantity
    frame_row5 = tk.Frame(frame_data_entry, bg="#F8F9FA")
    frame_row5.pack(fill=tk.X, pady=4)
    tk.Label(frame_row5, text="Total Quantity :", **lbl_style).pack(side=tk.LEFT)
    kamus_input["qty"] = tk.Entry(frame_row5, **ent_style)
    kamus_input["qty"].pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=3)

    # Field 6: Packing Quantity (Pcs)
    frame_row_pack = tk.Frame(frame_data_entry, bg="#F8F9FA")
    frame_row_pack.pack(fill=tk.X, pady=4)
    tk.Label(frame_row_pack, text="Packing Qty (Pcs) :", **lbl_style).pack(side=tk.LEFT)
    kamus_input["packing_qty"] = tk.Entry(frame_row_pack, **ent_style)
    kamus_input["packing_qty"].pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=3)
    
    # Field 7: Manufactured Date
    frame_row6 = tk.Frame(frame_data_entry, bg="#F8F9FA")
    frame_row6.pack(fill=tk.X, pady=4)
    tk.Label(frame_row6, text="Manufactured Date :", **lbl_style).pack(side=tk.LEFT)
    kamus_input["mfg"] = DateEntry(frame_row6, width=30, font=("Segoe UI", 10), background="#0D6EFD", foreground="white", borderwidth=1, date_pattern="dd/mm/yyyy")
    kamus_input["mfg"].pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=1)
    
    # Field 8: Machine Code (Dynamic Matching with Part Number)
    frame_row7 = tk.Frame(frame_data_entry, bg="#F8F9FA")
    frame_row7.pack(fill=tk.X, pady=4)
    tk.Label(frame_row7, text="Machine Code :", **lbl_style).pack(side=tk.LEFT)
    frame_machine_mix = tk.Frame(frame_row7, bg="#F8F9FA")
    frame_machine_mix.pack(side=tk.LEFT, fill=tk.X, expand=True)
    kamus_input["machine"] = tk.Entry(frame_machine_mix, font=("Segoe UI", 10), relief="groove", bd=1)
    kamus_input["machine"].pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=1)
    
    kamus_input["machine"].bind("<KeyRelease>", lambda e: cd.kendalikan_taipan_dropdown(
        kamus_input["machine"], win_inner, cd.ambil_machine_terikat(kamus_input["part"].get())
    ))
    
    tk.Button(frame_machine_mix, text="▼", font=("Segoe UI", 7), bg="#E2E8F0", fg="#475569", relief="groove", bd=1, width=3, takefocus=False,
              command=lambda: cd.papar_senarai_toplevel(kamus_input["machine"], win_inner, cd.ambil_machine_terikat(kamus_input["part"].get()))).pack(side=tk.RIGHT, fill=tk.Y)
    
    # Field 9: Lot Number
    frame_row8 = tk.Frame(frame_data_entry, bg="#F8F9FA")
    frame_row8.pack(fill=tk.X, pady=4)
    tk.Label(frame_row8, text="Lot Number :", **lbl_style).pack(side=tk.LEFT)
    kamus_input["lot"] = tk.Entry(frame_row8, **ent_style)
    kamus_input["lot"].pack(side=tk.LEFT, fill=tk.X, expand=True, ipady=3)

    # 🔗 PENGURUSAN IKATAN JUMP UTAMA (NO .ENTRY BUG)
    kamus_input["date"].bind("<Return>", lambda e: alih_fokus(e, kamus_input["customer"]))
    kamus_input["customer"].bind("<Return>", lambda e: alih_fokus(e, kamus_input["drawing"]))
    kamus_input["drawing"].bind("<Return>", lambda e: alih_fokus(e, kamus_input["part"]))
    kamus_input["part"].bind("<Return>", lambda e: alih_fokus(e, kamus_input["qty"]))
    kamus_input["qty"].bind("<Return>", lambda e: alih_fokus(e, kamus_input["packing_qty"]))
    kamus_input["packing_qty"].bind("<Return>", lambda e: alih_fokus(e, kamus_input["mfg"]))
    kamus_input["mfg"].bind("<Return>", lambda e: alih_fokus(e, kamus_input["machine"]))
    kamus_input["machine"].bind("<Return>", lambda e: alih_fokus(e, kamus_input["lot"]))
    kamus_input["lot"].bind("<Return>", lambda e: backend.proses_submit_data_pukal(win_inner, kamus_input))

    # ------------------ PANEL ACTIONS BUTTONS ------------------
    frame_butang_utama = tk.Frame(win_inner, bg="#F8F9FA")
    frame_butang_utama.pack(fill=tk.X, padx=30, pady=(0, 5))
    
    btn_submit = tk.Button(
        frame_butang_utama, text="SUBMIT & GENERATE BATCH INNER QR", 
        font=("Segoe UI", 10, "bold"), bg="#198754", fg="white", relief="flat", pady=8,
        command=lambda: backend.proses_submit_data_pukal(win_inner, kamus_input)
    )
    btn_submit.pack(fill=tk.X, pady=5)
    
    frame_butang_bawah = tk.Frame(win_inner, bg="#F8F9FA")
    frame_butang_bawah.pack(fill=tk.X, padx=30, pady=(0, 15))
    
    btn_import_prod = tk.Button(
        frame_butang_bawah, text="📥 EXCEL CUS", font=("Segoe UI", 8, "bold"), 
        bg="#0D6EFD", fg="white", relief="flat", pady=6,
        command=lambda: laksanakan_import_excel_langsung(win_inner)
    )
    btn_import_prod.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 2))
    
    btn_import_mach = tk.Button(
        frame_butang_bawah, text="📥 EXCEL MACH", font=("Segoe UI", 8, "bold"), 
        bg="#6F42C1", fg="white", relief="flat", pady=6,
        command=lambda: laksanakan_import_excel_machine(win_inner)
    )
    btn_import_mach.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=2)
    
    def laksanakan_reset():
        clear_isian_bawah()
        kamus_input["customer"].delete(0, tk.END)
        kamus_input["qty"].delete(0, tk.END)
        kamus_input["packing_qty"].delete(0, tk.END)
        kamus_input["machine"].delete(0, tk.END)
        kamus_input["lot"].delete(0, tk.END)
        
    btn_reset = tk.Button(
        frame_butang_bawah, text="🔄 RESET", font=("Segoe UI", 8, "bold"), 
        bg="#FD7E14", fg="white", relief="flat", pady=6, command=laksanakan_reset
    )
    btn_reset.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=2)
    
    btn_back = tk.Button(
        frame_butang_bawah, text="◀ BACK", font=("Segoe UI", 8, "bold"), 
        bg="#343A40", fg="white", relief="flat", pady=6, command=win_inner.destroy
    )
    btn_back.pack(side=tk.RIGHT, fill=tk.X, expand=True, padx=(2, 0))
    
    return win_inner
